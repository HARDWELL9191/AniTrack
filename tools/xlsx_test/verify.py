"""Opens the sample written by main.swift and checks it the way Excel would read it."""
import sys, zipfile, datetime
import xml.dom.minidom
from openpyxl import load_workbook

path = sys.argv[1]
z = zipfile.ZipFile(path)
assert z.testzip() is None, "bad CRC"
for name in z.namelist():
    if name.endswith((".xml", ".rels")):
        xml.dom.minidom.parseString(z.read(name))   # well-formed XML
print("zip + xml ok:", z.namelist())

wb = load_workbook(path)
names = wb.sheetnames
print("sheets:", names)
assert names[0] == "Watching"
assert names[1] == "BadNamex", names[1]
assert names[2] == "watching (2)", names[2]
assert len(names[3]) == 31
ws = wb["Watching"]
assert [c.value for c in ws[1]] == ["Title", "Other", "Score", "Progress", "Episodes", "Added", "AniList"]
assert ws["A1"].font.b, "header not bold"
assert ws.freeze_panes == "A2", ws.freeze_panes
assert ws.auto_filter.ref == "A1:G1005", ws.auto_filter.ref
assert ws["A2"].value == "Frieren: Beyond Journey's End"
assert ws["C2"].value == 10 and ws["D2"].value == 28
assert isinstance(ws["F2"].value, datetime.datetime), type(ws["F2"].value)
assert ws["G2"].value == '=HYPERLINK("https://anilist.co/anime/154587")', ws["G2"].value
assert ws["A3"].value == "Tom & Jerry <\"quoted\"> 'x'"
assert ws["B3"].value == "Čšž đ — 日本語 🎉"
assert ws["C3"].value == 8.5 and ws["C3"].number_format == "0.0"
assert ws["E3"].value is None and ws["F3"].value is None
assert ws["G3"].value == '=HYPERLINK("https://example.com/?a=""b""&c=1")', ws["G3"].value
assert ws["A4"].value == "  spaced  "
assert ws["B4"].value == "badchars", repr(ws["B4"].value)
assert ws["D4"].value == -2 and ws["E4"].value == 1234567 and ws["G4"].value is None
assert len(ws["A5"].value) == 32000
assert ws["A1005"].value == "Show 999"
assert ws.max_row == 1005
assert wb["BadNamex"].max_row == 1
print("all checks passed")
